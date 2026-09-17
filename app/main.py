"""Application FastAPI : dashboard de gestion des 101 comptes Instagram."""
import csv
import io
import re
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
from sqlalchemy import func
from starlette.middleware.sessions import SessionMiddleware

from .config import settings
from .database import (Account, AuditLog, InsightSnapshot, Post, PostTarget,
                       SessionLocal, User, get_db, init_db)
from .departments import DEPARTMENTS
from . import instagram_api as ig
from .scheduler import start_scheduler
from .security import (audit, can_access_account, encrypt_token,
                       hash_password, verify_password)

app = FastAPI(title="IG Manager - 101 departements")
app.add_middleware(SessionMiddleware, secret_key=settings.session_secret,
                   https_only=settings.https_only, same_site="lax")  # HTTPS_ONLY=1 dans .env en production
templates = Jinja2Templates(directory=str(Path(__file__).parent.parent / "templates"))
app.mount("/static", StaticFiles(directory=str(Path(__file__).parent.parent / "static")),
          name="static")


@app.on_event("startup")
def startup():
    init_db()
    db = SessionLocal()
    try:
        if db.query(Account).count() == 0:
            for code, name in DEPARTMENTS:
                db.add(Account(department_code=code, department_name=name))
            db.commit()
        if db.query(User).count() == 0:
            db.add(User(email=settings.admin_email,
                        password_hash=hash_password(settings.admin_password),
                        role="admin"))
            db.commit()
    finally:
        db.close()
    start_scheduler()


# ---------- Heures : saisie en heure de Paris, stockage/comparaison en UTC ----------

PARIS = ZoneInfo("Europe/Paris")
UTC = ZoneInfo("UTC")


def to_utc_naive(dt: datetime | None) -> datetime | None:
    """Heure saisie (Paris, sans fuseau) -> UTC sans fuseau, comme le planificateur."""
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=PARIS)
    return dt.astimezone(UTC).replace(tzinfo=None)


def to_paris_iso(dt: datetime | None) -> str | None:
    """UTC sans fuseau (base) -> 'AAAA-MM-JJTHH:MM' en heure de Paris."""
    if dt is None:
        return None
    return dt.replace(tzinfo=UTC).astimezone(PARIS).strftime("%Y-%m-%dT%H:%M")


# ---------- Auth ----------

def current_user(request: Request, db=Depends(get_db)) -> User:
    uid = request.session.get("user_id")
    if not uid:
        raise HTTPException(401, "Non connecte")
    user = db.get(User, uid)
    if not user or not user.is_active:
        raise HTTPException(401, "Compte desactive")
    return user


def require_admin(user: User = Depends(current_user)) -> User:
    if user.role != "admin":
        raise HTTPException(403, "Reserve a l'equipe nationale")
    return user


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return templates.TemplateResponse(request, "login.html")


@app.post("/login")
async def login(request: Request, db=Depends(get_db)):
    form = dict(await request.form())
    user = db.query(User).filter(User.email == form.get("email", "")).first()
    if not user or not verify_password(form.get("password", ""), user.password_hash):
        audit(db, form.get("email", "?"), "login_failed",
              ip=request.client.host if request.client else None)
        raise HTTPException(401, "Identifiants invalides")
    request.session["user_id"] = user.id
    audit(db, user.email, "login", ip=request.client.host if request.client else None)
    return RedirectResponse("/", status_code=303)


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


# ---------- Dashboard ----------

@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, db=Depends(get_db)):
    if not request.session.get("user_id"):
        return RedirectResponse("/login", status_code=303)
    user = db.get(User, request.session["user_id"])
    accounts = db.query(Account).order_by(Account.department_code).all()
    if user.role != "admin":
        accounts = [a for a in accounts if a.department_code == user.department_code]
    return templates.TemplateResponse(request, "dashboard.html",
                                      {"user": user, "accounts": accounts})


# ---------- Comptes ----------

class TokenIn(BaseModel):
    ig_user_id: str = ""
    ig_username: str = ""
    access_token: str = ""
    fb_page_id: str = ""
    fb_page_token: str = ""


@app.get("/api/accounts")
def list_accounts(user: User = Depends(current_user), db=Depends(get_db)):
    q = db.query(Account).order_by(Account.department_code)
    accounts = q.all() if user.role == "admin" else \
        q.filter(Account.department_code == user.department_code).all()
    return [{"code": a.department_code, "name": a.department_name,
             "username": a.ig_username, "status": a.status,
             "token_expires_at": a.token_expires_at} for a in accounts]


def _link_account(db, user, code: str, ig_user_id: str, ig_username: str,
                  access_token: str, fb_page_id: str = "",
                  fb_page_token: str = "") -> str | None:
    """Lie un compte (Instagram et/ou Page Facebook). None si OK, sinon erreur."""
    acc = db.query(Account).filter(Account.department_code == code).first()
    if not acc:
        return f"Departement inconnu : {code}"
    if not can_access_account(user, acc):
        return f"Acces refuse au departement {code}"
    # Un token Instagram suffit : l'identifiant et le nom sont demandes a Meta.
    if access_token and (not ig_user_id or not ig_username):
        try:
            me = ig.whoami_sync(access_token)
        except ig.InstagramAPIError as e:
            return f"{code} : token Instagram refuse par Meta ({e})"
        except Exception as e:  # reseau, JSON...
            return f"{code} : impossible de joindre Meta ({e})"
        ig_user_id = ig_user_id or me["user_id"]
        ig_username = ig_username or me["username"]
        if not ig_user_id:
            return f"{code} : Meta n'a pas renvoye d'identifiant pour ce token"
    has_ig = bool(ig_user_id and access_token)
    has_fb = bool(fb_page_id and fb_page_token)
    if not (has_ig or has_fb):
        return (f"{code} : fournir le token Instagram, "
                "ou fb_page_id + fb_page_token (Facebook)")
    if has_ig:
        acc.ig_user_id = ig_user_id
        acc.ig_username = ig_username
        acc.access_token_enc = encrypt_token(access_token)
        # token longue durée : 60 jours ; le planificateur le rafraîchit 10 jours avant
        acc.token_expires_at = datetime.utcnow() + timedelta(days=60)
    if has_fb:
        acc.fb_page_id = fb_page_id
        acc.fb_page_token_enc = encrypt_token(fb_page_token)
    acc.status = "active"
    db.commit()
    audit(db, user.email, "token_update",
          f"dept={code} ig={has_ig} fb={has_fb}")
    return None


@app.post("/api/accounts/{code}/token")
def set_token(code: str, body: TokenIn, user: User = Depends(current_user),
              db=Depends(get_db)):
    err = _link_account(db, user, code, body.ig_user_id, body.ig_username,
                        body.access_token, body.fb_page_id, body.fb_page_token)
    if err:
        raise HTTPException(404 if "inconnu" in err else
                            (403 if "refuse au" in err else 400), err)
    acc = db.query(Account).filter(Account.department_code == code).first()
    return {"ok": True, "username": acc.ig_username, "ig_user_id": acc.ig_user_id}


@app.post("/api/accounts/import")
async def import_accounts(file: UploadFile = File(...),
                          user: User = Depends(require_admin), db=Depends(get_db)):
    """Import en masse CSV ou XLSX.

    Colonnes : departement, ig_username, ig_user_id, access_token.
    """
    name = (file.filename or "").lower()
    data = await file.read()
    rows: list[dict] = []
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig")
        first = text.splitlines()[0] if text.splitlines() else ""
        delim = ";" if first.count(";") > first.count(",") else ","
        rows = list(csv.DictReader(io.StringIO(text), delimiter=delim))
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True).active
        it = ws.iter_rows(values_only=True)
        headers = [str(h or "").strip().lower() for h in next(it)]
        for r in it:
            rows.append({headers[i]: ("" if v is None else str(v).strip())
                         for i, v in enumerate(r) if i < len(headers)})
    else:
        raise HTTPException(400, "Format non supporte : utilisez .csv ou .xlsx")

    results = {"ok": 0, "ignores": 0, "erreurs": []}
    for row in rows:
        norm = {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        code = norm.get("departement") or norm.get("department_code") or norm.get("code")
        token = norm.get("access_token") or norm.get("token")
        ig_id = norm.get("ig_user_id") or norm.get("instagram_id")
        ig_name = norm.get("ig_username") or norm.get("username")
        fb_id = norm.get("fb_page_id") or norm.get("page_id")
        fb_tok = norm.get("fb_page_token") or norm.get("page_token")
        if not code or not (token or ig_id or fb_tok):
            results["ignores"] += 1
            continue
        err = _link_account(db, user, code, ig_id, ig_name, token, fb_id, fb_tok)
        if err:
            results["erreurs"].append(err)
        else:
            results["ok"] += 1
    audit(db, user.email, "bulk_import",
          f"ok={results['ok']} erreurs={len(results['erreurs'])}")
    return results


@app.post("/api/accounts/{code}/suspend")
def suspend_account(code: str, user: User = Depends(require_admin),
                    db=Depends(get_db)):
    """KILL SWITCH : gele un compte (plus aucune publication via l'app)."""
    acc = db.query(Account).filter(Account.department_code == code).first()
    if not acc:
        raise HTTPException(404)
    acc.status = "suspended"
    acc.access_token_enc = None  # revocation locale immediate
    db.commit()
    audit(db, user.email, "account_suspend", f"dept={code}")
    return {"ok": True, "message": "Compte gele. Revoquez aussi le token cote Meta."}


# ---------- Publication ----------

class PostIn(BaseModel):
    caption: str
    media_url: str
    media_type: str = "IMAGE"
    platform: str = "instagram"        # instagram | facebook | both
    scheduled_at: datetime | None = None
    targets: list[str] | str = "all"   # "all" ou liste de codes departement
    label: str = ""                    # libellé du plan / de la campagne (optionnel)


def _check_post(db, user: User, body: PostIn):
    """Valide un envoi (droits, cibles, URLs, format). Renvoie (comptes, urls, plateformes)."""
    if body.targets == "all" and user.role != "admin":
        raise HTTPException(403, "Seule l'equipe nationale peut publier sur tous les comptes")
    codes = [c for c, _ in DEPARTMENTS] if body.targets == "all" else list(body.targets)
    if not codes:
        raise HTTPException(400, "Aucune cible : mettre all ou des codes de departement")
    accounts = db.query(Account).filter(Account.department_code.in_(codes)).all()
    unknown = set(codes) - {a.department_code for a in accounts}
    if unknown:
        raise HTTPException(400, f"Departement(s) inconnu(s) : {', '.join(sorted(unknown))}")
    for acc in accounts:
        if not can_access_account(user, acc):
            raise HTTPException(403, f"Acces refuse au departement {acc.department_code}")
    if not body.media_url.strip():
        raise HTTPException(400, "L'URL du media est obligatoire")
    for u in body.media_url.replace("|", "\n").splitlines():
        u = u.strip()
        if u and not (u.startswith("http://") or u.startswith("https://")):
            raise HTTPException(400, f"URL de media non conforme : {u[:80]}")
    if body.platform not in ("instagram", "facebook", "both"):
        raise HTTPException(400, "platform doit etre instagram, facebook ou both")
    urls = [u.strip() for u in body.media_url.replace("|", "\n").splitlines() if u.strip()]
    if body.media_type == "IMAGE" and len(urls) > 1:
        body.media_type = "CAROUSEL"   # bascule automatique
    if body.media_type == "CAROUSEL" and not 2 <= len(urls) <= 10:
        raise HTTPException(400, "Un carrousel demande 2 a 10 images")
    if body.media_type in ("REELS", "STORIES") and len(urls) != 1:
        raise HTTPException(400, "Un reel ou une story demande une seule URL")
    platforms = ["instagram", "facebook"] if body.platform == "both" else [body.platform]
    return accounts, urls, platforms


def _create_post(db, user: User, body: PostIn) -> tuple[Post, int]:
    """Cree un post + ses cibles. Utilise par /api/posts et par l'import de plan."""
    accounts, urls, platforms = _check_post(db, user, body)
    post = Post(caption=body.caption, media_url=body.media_url,
                media_type=body.media_type, platform=body.platform,
                scheduled_at=to_utc_naive(body.scheduled_at),
                label=(body.label or "").strip()[:80],
                created_by=user.id, is_national=(body.targets == "all"))
    db.add(post)
    db.flush()
    for acc in accounts:
        for p in platforms:
            db.add(PostTarget(post_id=post.id, account_id=acc.id, platform=p))
    db.commit()
    audit(db, user.email, "publish",
          f"post={post.id} targets={len(accounts)} national={post.is_national}")
    return post, len(accounts) * len(platforms)


@app.post("/api/posts")
def create_post(body: PostIn, user: User = Depends(current_user), db=Depends(get_db)):
    post, n = _create_post(db, user, body)
    return {"post_id": post.id, "targets": n,
            "mode": "planifie" if post.scheduled_at else "publication sous 1 min"}


@app.get("/api/posts/{post_id}/status")
def post_status(post_id: int, user: User = Depends(current_user), db=Depends(get_db)):
    targets = db.query(PostTarget).filter(PostTarget.post_id == post_id).all()
    return [{"dept": t.account.department_code, "platform": t.platform,
             "status": t.status, "error": t.error} for t in targets]


# ---------- Planification ----------

_TYPE_MAP = {"post": "IMAGE", "photo": "IMAGE", "image": "IMAGE", "": "IMAGE",
             "carrousel": "CAROUSEL", "carousel": "CAROUSEL",
             "reel": "REELS", "reels": "REELS", "video": "REELS", "vidéo": "REELS",
             "story": "STORIES", "stories": "STORIES"}
_TYPE_LABEL = {"IMAGE": "post", "CAROUSEL": "carrousel", "REELS": "reel", "STORIES": "story"}
_DATE_FORMATS = ("%d/%m/%Y %H:%M", "%d/%m/%Y %Hh%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%y %H:%M",
                 "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M:%S")


def _parse_when(value) -> datetime | None:
    """Heure de Paris -> datetime sans fuseau. Vide = immediat."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    s = str(value).strip()
    if s.lower() in ("", "maintenant", "now", "immediat", "immédiat"):
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    raise ValueError(f"date illisible « {s} » (attendu : 18/09/2026 09:30)")


def _parse_targets(value) -> list[str] | str:
    s = str(value or "").strip().lower()
    if s in ("all", "tous", "tout", "*", "national"):
        return "all"
    codes = [c for c in re.split(r"[,\s|/;]+", s) if c]
    if not codes:
        raise ValueError("cibles vides (mettre all, ou des codes : 75, 13, 69)")
    known = {c for c, _ in DEPARTMENTS}
    out = []
    for c in codes:
        c = c.upper()
        if c.isdigit() and len(c) == 1:
            c = "0" + c
        if c not in known:
            raise ValueError(f"departement inconnu : {c}")
        if c not in out:
            out.append(c)
    return out


def _read_table(name: str, data: bytes) -> list[dict]:
    """CSV (virgule ou point-virgule) ou XLSX -> liste de dict, en-tetes en minuscules."""
    name = (name or "").lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig")
        first = text.splitlines()[0] if text.splitlines() else ""
        delim = ";" if first.count(";") > first.count(",") else ","
        return [{(k or "").strip().lower(): (v if v is not None else "") for k, v in r.items()}
                for r in csv.DictReader(io.StringIO(text), delimiter=delim)]
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        headers = [str(h or "").strip().lower() for h in next(it, [])]
        rows = []
        for r in it:
            if r is None or all(v in (None, "") for v in r):
                continue
            rows.append({headers[i]: (v if v is not None else "")
                         for i, v in enumerate(r) if i < len(headers)})
        return rows
    raise HTTPException(400, "Format non supporte : utilisez .csv ou .xlsx")


def _serialize_post(p: Post) -> dict:
    counts = {"pending": 0, "published": 0, "failed": 0, "rate_limited": 0, "cancelled": 0}
    for t in p.targets:
        counts[t.status] = counts.get(t.status, 0) + 1
    total = len(p.targets)
    now = datetime.utcnow()
    if total and counts["published"] == total:
        state = "publié"
    elif counts["pending"]:
        state = "programmé" if (p.scheduled_at and p.scheduled_at > now) else "en cours"
    elif counts["failed"] or counts["rate_limited"]:
        state = "partiel" if counts["published"] else "échecs"
    elif total and counts["cancelled"] == total:
        state = "annulé"
    else:
        state = "terminé"
    urls = [u.strip() for u in (p.media_url or "").replace("|", "\n").splitlines() if u.strip()]
    return {"id": p.id, "label": p.label or "", "when": to_paris_iso(p.scheduled_at),
            "created_at": to_paris_iso(p.created_at),
            "media_type": p.media_type, "type_label": _TYPE_LABEL.get(p.media_type, p.media_type),
            "platform": p.platform, "media_url": p.media_url or "",
            "first_url": urls[0] if urls else "", "n_urls": len(urls),
            "caption": p.caption or "", "is_national": bool(p.is_national),
            "total": total, "counts": counts, "state": state,
            "targets": sorted({t.account.department_code for t in p.targets})}


@app.get("/api/plan")
def plan_list(user: User = Depends(current_user), db=Depends(get_db),
              include_done: bool = False, days: int = 30):
    """Le plan : envois a venir (et termines si include_done)."""
    since = datetime.utcnow() - timedelta(days=days)
    posts = db.query(Post).filter((Post.created_at >= since) | (Post.scheduled_at >= since)).all()
    if user.role != "admin":
        posts = [p for p in posts
                 if any(t.account.department_code == user.department_code for t in p.targets)]
    out = [_serialize_post(p) for p in posts]
    if not include_done:
        out = [o for o in out if o["counts"]["pending"] or o["counts"]["failed"]
               or o["counts"]["rate_limited"]]
    out.sort(key=lambda o: (o["when"] or o["created_at"] or ""))
    return out


@app.post("/api/posts/{post_id}/cancel")
def cancel_post(post_id: int, user: User = Depends(current_user), db=Depends(get_db)):
    """Annule les cibles encore en attente d'un envoi."""
    post = db.get(Post, post_id)
    if not post:
        raise HTTPException(404, "Envoi inconnu")
    n = 0
    for t in post.targets:
        if t.status == "pending" and can_access_account(user, t.account):
            t.status, t.error = "cancelled", "annule avant envoi"
            n += 1
    db.commit()
    audit(db, user.email, "post_cancel", f"post={post_id} cibles={n}")
    return {"cancelled": n}


@app.post("/api/posts/{post_id}/retry")
def retry_post(post_id: int, user: User = Depends(current_user), db=Depends(get_db)):
    """Remet en attente les cibles en echec (relance sous 1 min)."""
    post = db.get(Post, post_id)
    if not post:
        raise HTTPException(404, "Envoi inconnu")
    n = 0
    for t in post.targets:
        if t.status in ("failed", "rate_limited") and can_access_account(user, t.account):
            t.status, t.error = "pending", None
            n += 1
    db.commit()
    audit(db, user.email, "post_retry", f"post={post_id} cibles={n}")
    return {"retried": n}


@app.post("/api/plan/import")
async def plan_import(file: UploadFile = File(...), dry_run: bool = False,
                      user: User = Depends(require_admin), db=Depends(get_db)):
    """Import d'un plan : une ligne = un envoi.

    Colonnes : quand, cibles, type, visuel, legende, libelle, plateforme.
    dry_run=true : relit et valide le plan sans rien programmer.
    """
    rows = _read_table(file.filename, await file.read())
    lines, errors, created = [], [], 0

    def g(row, *keys):
        for k in keys:
            v = row.get(k)
            if v not in (None, ""):
                return str(v).strip()
        return ""

    for i, row in enumerate(rows, start=2):   # ligne 1 = en-tete
        visuel = g(row, "visuel", "visuels", "image", "media", "url", "media_url")
        legende = g(row, "legende", "légende", "caption", "texte")
        cibles = g(row, "cibles", "comptes", "departements", "départements")
        if not (visuel or legende or cibles):
            continue  # ligne vide
        entry = {"ligne": i, "quand": "", "cibles": "", "n_cibles": 0, "type": "",
                 "visuel": visuel.replace("|", "\n"), "legende": legende[:2200],
                 "libelle": g(row, "libelle", "libellé", "label", "campagne")[:80], "statut": "ok"}
        try:
            when = _parse_when(row.get("quand") if row.get("quand") not in (None, "")
                               else g(row, "date", "heure"))
            targets = _parse_targets(cibles)
            t = g(row, "type", "format").lower()
            if t not in _TYPE_MAP:
                raise ValueError(f"type inconnu « {t} » (post, carrousel, reel, story)")
            plat = (g(row, "plateforme", "platform") or "instagram").lower()
            plat = {"les deux": "both", "deux": "both"}.get(plat, plat)
            if plat not in ("instagram", "facebook", "both"):
                raise ValueError(f"plateforme inconnue « {plat} »")
            if not visuel:
                raise ValueError("visuel manquant (URL https publique)")
            body = PostIn(caption=entry["legende"], media_url=entry["visuel"],
                          media_type=_TYPE_MAP[t], platform=plat, scheduled_at=when,
                          targets=targets, label=entry["libelle"])
            accounts, urls, platforms = _check_post(db, user, body)   # meme validation que l'envoi
            entry["quand"] = when.strftime("%d/%m/%Y %H:%M") if when else "immédiat"
            entry["cibles"] = "tous" if targets == "all" else ", ".join(targets)
            entry["n_cibles"] = len(accounts) * len(platforms)
            entry["type"] = _TYPE_LABEL[body.media_type] + (f" ({len(urls)})" if len(urls) > 1 else "")
            if not dry_run:
                post, _ = _create_post(db, user, body)
                entry["post_id"] = post.id
                created += 1
        except HTTPException as e:
            entry["statut"] = f"erreur : {e.detail}"
            errors.append(f"ligne {i} : {e.detail}")
        except ValueError as e:
            entry["statut"] = f"erreur : {e}"
            errors.append(f"ligne {i} : {e}")
        lines.append(entry)
    if not dry_run:
        audit(db, user.email, "plan_import",
              f"fichier={file.filename} envois={created} erreurs={len(errors)}")
    ok = created if not dry_run else sum(1 for l in lines if l["statut"] == "ok")
    return {"dry_run": dry_run, "ok": ok, "erreurs": errors, "lignes": lines}


# ---------- Statistiques ----------

@app.get("/api/stats/overview")
def stats_overview(user: User = Depends(current_user), db=Depends(get_db)):
    """Vue nationale : totaux et top departements (dernier snapshot)."""
    latest = (db.query(InsightSnapshot.account_id,
                       func.max(InsightSnapshot.date).label("d"))
              .group_by(InsightSnapshot.account_id).subquery())
    rows = (db.query(InsightSnapshot).join(
        latest, (InsightSnapshot.account_id == latest.c.account_id) &
                (InsightSnapshot.date == latest.c.d)).all())
    total_followers = sum(r.followers for r in rows)
    total_reach = sum(r.reach for r in rows)
    top = sorted(rows, key=lambda r: r.followers, reverse=True)[:10]
    active = db.query(Account).filter(Account.status == "active").count()
    return {"accounts_reporting": len(rows),
            "active_accounts": active,
            "total_followers": total_followers,
            "total_reach_yesterday": total_reach,
            "top_departments": [{"dept": r.account.department_code,
                                 "name": r.account.department_name,
                                 "followers": r.followers} for r in top]}


@app.get("/api/stats/{code}")
def stats_department(code: str, user: User = Depends(current_user), db=Depends(get_db)):
    acc = db.query(Account).filter(Account.department_code == code).first()
    if not acc:
        raise HTTPException(404)
    if not can_access_account(user, acc):
        raise HTTPException(403)
    snaps = (db.query(InsightSnapshot).filter(InsightSnapshot.account_id == acc.id)
             .order_by(InsightSnapshot.date).all())
    return [{"date": s.date, "followers": s.followers, "reach": s.reach,
             "profile_views": s.profile_views} for s in snaps]


# ---------- Audit (admin) ----------

@app.get("/api/audit")
def audit_log(user: User = Depends(require_admin), db=Depends(get_db), limit: int = 200):
    rows = db.query(AuditLog).order_by(AuditLog.timestamp.desc()).limit(limit).all()
    return [{"ts": r.timestamp, "user": r.user_email, "action": r.action,
             "detail": r.detail, "ip": r.ip} for r in rows]


# ---------- Utilisateurs (admin) ----------

class UserIn(BaseModel):
    email: str
    password: str
    department_code: str


@app.post("/api/users")
def create_user(body: UserIn, user: User = Depends(require_admin), db=Depends(get_db)):
    """Cree un correspondant departemental (acces limite a son departement)."""
    if db.query(User).filter(User.email == body.email).first():
        raise HTTPException(409, "Email deja utilise")
    db.add(User(email=body.email, password_hash=hash_password(body.password),
                role="correspondent", department_code=body.department_code))
    db.commit()
    audit(db, user.email, "user_create", f"{body.email} dept={body.department_code}")
    return {"ok": True}


@app.post("/api/users/{email}/deactivate")
def deactivate_user(email: str, user: User = Depends(require_admin), db=Depends(get_db)):
    """Desactive un correspondant (depart, compromission) - effet immediat."""
    target = db.query(User).filter(User.email == email).first()
    if not target:
        raise HTTPException(404)
    target.is_active = False
    db.commit()
    audit(db, user.email, "user_deactivate", email)
    return {"ok": True}
